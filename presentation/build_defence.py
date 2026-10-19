# -*- coding: utf-8 -*-
"""Сборка презентации к защите паспорта на шаблоне курса.

Источник контента — заполненный паспорт методической разработки
(занятие «Цифровая редакция школы»), дизайн — «Шаблон презентации к защите.pptx».

Запуск:  .venv\\Scripts\\python.exe presentation\\build_defence.py
Выход:   <корень проекта>\\Презентация-к-защите-Цифровая-редакция.pptx
"""
from copy import deepcopy
from pathlib import Path

import qrcode
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent                 # school-portal/presentation
ROOT = HERE.parents[1]                                 # «Проект по умолчанию»
TEMPLATE = ROOT / "Входящие" / "Шаблон презентации к защите.pptx"
OUT = ROOT / "Презентация-к-защите-Цифровая-редакция.pptx"
LOGO = HERE / "logo.png"

PORTAL = "https://school-portal-61--sochinart1.replit.app"
LANDING = "https://annasochin.github.io/school-portal-61/"
GITHUB = "https://github.com/annasochin/school-portal-61"

TOL = 0.04  # дюйма — допуск при поиску формы по координатам


# ---------------------------------------------------------------- утилиты

def set_text(shape, lines, link=None):
    """Заменяет текст формы, сохраняя шрифт/цвет/выравнивание исходного шаблона."""
    if isinstance(lines, str):
        lines = [lines]
    tf = shape.text_frame
    tx_body = tf._txBody

    p_tpl = r_tpl = None
    for p in tx_body.findall(qn("a:p")):
        if p_tpl is None:
            p_tpl = p.find(qn("a:pPr"))
        for r in p.findall(qn("a:r")):
            r_tpl = r.find(qn("a:rPr"))
            break
        if r_tpl is not None:
            break

    for p in tx_body.findall(qn("a:p")):
        tx_body.remove(p)

    for line in lines:
        p = etree.SubElement(tx_body, qn("a:p"))
        if p_tpl is not None:
            p.append(deepcopy(p_tpl))
        r = etree.SubElement(p, qn("a:r"))
        if r_tpl is not None:
            r.append(deepcopy(r_tpl))
        t = etree.SubElement(r, qn("a:t"))
        t.text = line

    if link:
        for p in tf.paragraphs:
            for run in p.runs:
                run.hyperlink.address = link


def make_round_logo():
    """Возвращает путь к логотипу с прозрачным угловой фоном (маска по кругу).

    Исходный logo.png залит белым — на тёмных слайдах он виден белым квадратом.
    """
    from PIL import Image, ImageDraw

    out = HERE / "logo_round.png"
    im = Image.open(LOGO).convert("RGB")
    w, h = im.size

    # границы эмблемы — всё, что не белое
    px = im.load()
    xs, ys = [], []
    step = max(1, w // 256)
    for y in range(0, h, step):
        for x in range(0, w, step):
            r, g, b = px[x, y]
            if r < 245 or g < 245 or b < 245:
                xs.append(x)
                ys.append(y)
    cx, cy = (min(xs) + max(xs)) / 2, (min(ys) + max(ys)) / 2
    rad = max(max(xs) - min(xs), max(ys) - min(ys)) / 2 + 3

    mask = Image.new("L", (w, h), 0)
    ImageDraw.Draw(mask).ellipse(
        (cx - rad, cy - rad, cx + rad, cy + rad), fill=255)
    rgba = im.convert("RGBA")
    rgba.putalpha(mask)
    rgba.save(out)
    return out


def find_shape(slide, left, top):
    """Ищет форму по координатам; при наложении — ту, где есть текст."""
    hits = []
    for sh in slide.shapes:
        if sh.left is None or sh.top is None:
            continue
        if abs(Inches(left).emu - sh.left) < Inches(TOL).emu and \
           abs(Inches(top).emu - sh.top) < Inches(TOL).emu:
            hits.append(sh)
    if len(hits) == 1:
        return hits[0]
    for sh in hits:                      # наложение слоёв: берём текстовый
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh
    return None


# ---------------------------------------------------------------- содержание
# (left, top) — координаты формы в шаблоне; текст — из паспорта занятия.

FILLS = {
    0: [  # обложка
        (0.7, 1.91, ["Паспорт методической разработки:",
                     "«Цифровая редакция школы»"]),
    ],
    1: [  # 1. Название занятия
        (0.7, 1.99,
         "Команды 7–9 классов публикуют настоящие анонсы олимпиад на школьном "
         "портале — от первого источника до печатной карточки с QR."),
        (0.98, 2.88,
         "«Цифровая редакция школы»: работающий анонс олимпиады "
         "на школьном портале"),
        (0.98, 4.7, "Информатика + русский + математика"),
        (5.03, 4.7, "7–9 классы (13–15 лет)"),
        (9.07, 4.7, "Компьютерный кабинет, 2 × 45 мин"),
        (0.98, 6.01, "Проверка формулировки пройдена:"),
        (0.98, 6.42,
         "Тема читается как заголовок: «Цифровая редакция школы» — ученики "
         "работают редакцией и публикуют анонсы олимпиад на портале школы. "
         "Понятно за 5 секунд без пояснений."),
    ],
    2: [  # 2. Проблема
        (0.7, 1.99,
         "Реальный вызов, на который отвечает занятие, и проектный вопрос."),
        (0.98, 2.88,
         "Объявления об олимпиадах теряются: чаты классов, бумажная доска, "
         "устные сообщения — единого места, где собраны все конкурсы, нет."),
        (0.98, 4.01,
         "Нельзя решить «в одно действие»: источники разные, сроки текучие, "
         "оформление у всех своё, а пропущенная дата — это пропущенная "
         "олимпиада. Страдают все: ученики, учителя, администрация."),
        (0.98, 5.18,
         "Проектный вопрос: как подготовить и опубликовать анонс реальной "
         "олимпиады, чтобы любой ученик за 30 секунд нашёл его, понял сроки "
         "и смог участвовать?"),
    ],
    3: [  # 3. Замысел (междисциплинарные связи)
        (1.42, 3.96, "Факты и сроки"),
        (1.48, 5.28, "Русский язык · математика"),
        (4.42, 3.86,
         "Портал собирает анонс из полей формы (тип, предмет, даты), "
         "публикует его и делает видимым всем: фильтр, поиск и календарь "
         "находят олимпиаду за секунды. Руками за то же время объявить всем "
         "не успеть — компьютер берёт на себя структуру, сроки и публикацию."),
    ],
    4: [  # 4. Этапы
        (0.7, 1.78,
         "Этапы 1–3 — занятие 1 (45 минут), этапы 4–5 — занятие 2 "
         "(45 минут): что делает команда на каждом этапе."),
        (1.8, 2.34, "Погружение и роли"),
        (5.63, 2.34, "Получают кейс заказчика, читают рабочий лист, делят 5 ролей"),
        (11.05, 2.38, "10 мин"),
        (1.8, 3.34, "Отбор источников и сверка дат"),
        (5.63, 3.34, "Находят 2–3 официальные олимпиады, сверяют сроки "
                     "с первоисточником"),
        (11.05, 3.39, "20 мин"),
        (1.8, 4.35, "Черновик анонса"),
        (5.63, 4.35, "Пишут «заголовок → суть → сроки → призыв», "
                     "вычитывают и сверяют даты"),
        (11.05, 4.39, "15 мин"),
        (1.8, 5.36, "Публикация и проверка"),
        (5.63, 5.36, "Заполняют форму, отправляют на модерацию, находят "
                     "анонс фильтром и поиском"),
        (11.05, 5.4, "20 мин"),
        (1.8, 6.31, "Карточка, QR и защита"),
        (5.63, 6.31, "Печатают карточку, проверяют QR, готовят рекомендацию "
                     "и защищают работу"),
        (11.05, 6.35, "25 мин"),
    ],
    5: [  # 5. Продукт и лендинг
        (0.7, 1.99,
         "Что появляется у учеников на выходе и где живёт лендинг занятия."),
        (1.27, 3.64,
         "Опубликованная страница реальной олимпиады на портале: точные сроки, "
         "описание, ссылка первоисточника. И печатная карточка с QR для доски "
         "объявлений — открыть, показать, перейти."),
        (7.41, 3.48, "🔗 annasochin.github.io/school-portal-61", LANDING),
        (7.47, 5.31, "Этапы занятия и предмет результата"),
        (7.47, 5.84, "Рабочий лист для печати и метрики оценки"),
    ],
    6: [  # 6. Как оценим
        (0.7, 1.99,
         "Четыре метрики занятия, рубрикатор защиты и чек-лист приёмки "
         "продукта."),
        (0.98, 3.37, "≥ 90 %"),
        (0.98, 4.17,
         "доля учеников, чей анонс опубликован и находится на главной через "
         "фильтр и поиск"),
        (7.47, 3.37, "≥ 80 % — команд с датами, совпадающими с первоисточником"),
        (7.47, 3.81, "≥ 2 — проверенных источника в среднем на команду"),
        (7.47, 4.24, "≥ 90 % — учеников, проверивших результат ИИ "
                     "по первоисточнику"),
        (7.32, 6.03, "Защита оценена по рубрикатору: 5 × 0–3 балла"),
    ],
    7: [  # 7. Куда дальше
        (0.98, 3.53,
         "Готовое к запуску для другого учителя: рабочий лист, чек-листы, "
         "инструкция; из оборудования — ПК и принтер. Учителю ~1 час: "
         "раздать листы, завести аккаунты, одобрить анонсы."),
        (5.27, 3.53,
         "Ученики разных параллелей как единая редакция: роли и чек-лист уже "
         "готовы; добавляются классные часы и общая доска объявлений."),
        (9.56, 3.53,
         "Портал как единая точка публикации олимпиад района: курсы для "
         "учителей, доступ к интернету и сети печати."),
        (0.98, 6.49,
         "Модерация задерживает публикацию — публикуем накануне; не у всех ПК "
         "— черновик готовят в паре; запасной вариант: печатная доска "
         "объявлений."),
    ],
    8: [  # финал
        (1.31, 2.95, "Спасибо за внимание!"),
        (1.31, 4.32,
         ["Портал: school-portal-61--sochinart1.replit.app",
          "Лендинг: annasochin.github.io/school-portal-61",
          "GitHub: github.com/annasochin/school-portal-61"]),
    ],
}

FINAL_LINKS = [PORTAL, LANDING, GITHUB]   # по порядку строк финального слайда


# ---------------------------------------------------------------- сборка

def main():
    prs = Presentation(str(TEMPLATE))
    problems = []

    for si, fills in FILLS.items():
        slide = prs.slides[si]
        for item in fills:
            left, top, text = item[0], item[1], item[2]
            link = item[3] if len(item) > 3 else None
            sh = find_shape(slide, left, top)
            if sh is None:
                problems.append(f"слайд {si + 1}: форма ({left}, {top}) не найдена")
                continue
            set_text(sh, text, link)

    # гиперссылки финального слайда
    fin = prs.slides[8]
    link_box = find_shape(fin, 1.31, 4.32)
    for para, url in zip(link_box.text_frame.paragraphs, FINAL_LINKS):
        for run in para.runs:
            run.hyperlink.address = url
            run.font.color.rgb = RGBColor(0x9E, 0xD9, 0xFF)  # по тёмному фону

    # подпись-подзаголовок обложки
    cover = prs.slides[0]
    sub = cover.shapes.add_textbox(Inches(0.7), Inches(3.75),
                                   Inches(11.9), Inches(1.5))
    tf = sub.text_frame
    tf.word_wrap = True
    lines = [
        "Занятие: работающий анонс олимпиады на школьном портале "
        "МБОУ СОШ № 61 им. П. Е. Ладана",
        "7–9 классы · 2 × 45 минут · компьютерный кабинет",
        "Программа «Информатика в междисциплинарных проектах: "
        "Архитектура образовательных решений»",
    ]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.size = Pt(17 if i == 0 else 14)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        run.font.bold = i == 0

    # логотип школы: обложка и финал — крупно, контентные слайды — справа
    # от заголовка раздела
    if LOGO.exists():
        logo = make_round_logo()
        cover.shapes.add_picture(str(logo), Inches(0.7), Inches(5.55),
                                 height=Inches(1.0))
        for i in range(1, 8):
            prs.slides[i].shapes.add_picture(str(logo), Inches(11.62),
                                             Inches(1.14), height=Inches(0.62))
        fin.shapes.add_picture(str(logo), Inches(1.31), Inches(5.75),
                               height=Inches(0.8))

    # QR с сайтом портала на финальном слайде
    qr_path = HERE / "qr_site.png"
    img = qrcode.make(PORTAL)
    img.save(qr_path)
    fin.shapes.add_picture(str(qr_path), Inches(11.4), Inches(5.5),
                           width=Inches(1.2))
    cap = fin.shapes.add_textbox(Inches(11.0), Inches(5.15),
                                 Inches(1.9), Inches(0.3))
    p = cap.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "Сайт с QR"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    prs.save(str(OUT))
    if problems:
        print("ПРОБЛЕМЫ:")
        for p_ in problems:
            print(" -", p_)
    else:
        print(f"OK: {OUT}")


if __name__ == "__main__":
    main()
