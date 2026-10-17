# -*- coding: utf-8 -*-
"""Сборка презентации портала школы № 61 (7 слайдов, PPTX).

Запуск:  .venv\\Scripts\\python.exe presentation\\build_pptx.py
Выход:   <корень проекта>\\Презентация-портал-61.pptx
"""
import io
from pathlib import Path

import qrcode
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent            # school-portal/presentation
ROOT = HERE.parents[1]                            # «Проект по умолчанию»
OUT = ROOT / "Презентация-портал-61.pptx"
LOGO = HERE.parent / "logo-school-61" / "logo-school-61-512.png"

SITE = "https://school-portal-61--sochinart1.replit.app"
GITHUB = "https://github.com/annasochin/school-portal-61"
SCHOOL = "МБОУ СОШ № 61 им. П. Е. Ладана"

# цвета школы: белый / синий / голубой
BLUE = RGBColor(0x25, 0x63, 0xEB)     # основной синий
DEEP = RGBColor(0x1E, 0x3A, 0x8A)     # тёмно-синий
SKY = RGBColor(0x38, 0xBD, 0xF8)      # голубой
LIGHT = RGBColor(0xEF, 0xF6, 0xFF)    # светло-голубой фон
BORDER = RGBColor(0xDB, 0xE4, 0xEE)   # рамки
TEXT = RGBColor(0x1F, 0x29, 0x37)     # текст
MUTED = RGBColor(0x6B, 0x72, 0x80)    # подписи
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Segoe UI"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------- утилиты

def rect(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is not None:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    else:
        sp.fill.background()
    if line is not None:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def put(tf, text, size=15, bold=False, color=TEXT, first=False,
        align=PP_ALIGN.LEFT, space_before=0, space_after=6, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = 1.08
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return p


def bullets(slide, x, y, w, h, items, size=15.5, gap=13):
    tf = textbox(slide, x, y, w, h)
    for i, item in enumerate(items):
        put(tf, "•  " + item, size=size, first=(i == 0), space_after=gap)
    return tf


def header(slide, title, n):
    """Синяя шапка с логотипом школы + подвал со счётчиком слайдов."""
    rect(slide, 0, 0, 13.333, 1.0, fill=BLUE)
    # белый кружок-подложка, чтобы логотип читался на синем фоне
    rect(slide, 0.5, 0.08, 0.84, 0.84, fill=WHITE, shape=MSO_SHAPE.OVAL)
    slide.shapes.add_picture(str(LOGO), Inches(0.54), Inches(0.12),
                             width=Inches(0.76), height=Inches(0.76))
    tf = textbox(slide, 1.55, 0.14, 11.2, 0.72, anchor=MSO_ANCHOR.MIDDLE)
    put(tf, title, size=26, bold=True, color=WHITE, first=True, space_after=0)
    rect(slide, 0.6, 7.0, 12.13, 0.012, fill=BORDER)
    f = textbox(slide, 0.6, 7.08, 9.5, 0.32)
    put(f, f"{SCHOOL} · {SITE}", size=9.5, color=MUTED, first=True, space_after=0)
    f2 = textbox(slide, 11.4, 7.08, 1.33, 0.32)
    put(f2, f"{n} / 7", size=9.5, color=MUTED, first=True, space_after=0,
        align=PP_ALIGN.RIGHT)


def new_slide(title, n):
    s = prs.slides.add_slide(BLANK)
    header(s, title, n)
    return s


def picture(slide, path, x, y, w):
    """Картинка фиксированной ширины, возвращает высоту (с учётом пропорций)."""
    img = Image.open(path)
    h = w * img.height / img.width
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    return h


def caption(slide, x, y, w, text):
    tf = textbox(slide, x, y, w, 0.3)
    put(tf, text, size=10.5, color=MUTED, first=True, space_after=0,
        align=PP_ALIGN.CENTER)


def shot(slide, name, x, y, w, cap):
    """Скриншот в светлой рамке + подпись."""
    img = Image.open(HERE / name)
    h = w * img.height / img.width
    rect(slide, x - 0.08, y - 0.08, w + 0.16, h + 0.16, fill=WHITE, line=BORDER)
    picture(slide, HERE / name, x, y, w)
    caption(slide, x, y + h + 0.12, w, cap)


def link_run(p, text, url, size=16):
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = BLUE
    run.hyperlink.address = url
    return run


# ---------------------------------------------------------------- слайд 1

s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, 13.333, 7.5, fill=BLUE)
rect(s, 1.4, 1.15, 10.53, 4.0, fill=WHITE, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
picture(s, LOGO, 2.1, 1.7, 2.9)
tf = textbox(s, 5.5, 1.62, 6.1, 3.1, anchor=MSO_ANCHOR.MIDDLE)
put(tf, "Портал олимпиад и конкурсов", size=34, bold=True, color=DEEP,
    first=True, space_after=8)
put(tf, SCHOOL, size=18, color=TEXT, space_after=10)
put(tf, "Учителя публикуют · ученики находят", size=14.5, color=MUTED, italic=True,
    space_after=0)
tf = textbox(s, 1.4, 5.5, 10.53, 1.2, anchor=MSO_ANCHOR.TOP)
put(tf, "Презентация школьного информационного портала · 2026", size=15,
    color=WHITE, first=True, align=PP_ALIGN.CENTER, space_after=6)
p = put(tf, "", size=17, color=SKY, align=PP_ALIGN.CENTER, space_after=0)
link_run(p, SITE, SITE, size=17)

# ---------------------------------------------------------------- слайд 2

s = new_slide("Замысел: одна точка входа для всей информации школы", 2)
bullets(s, 0.6, 1.4, 7.5, 5.3, [
    "«Одно окно» для олимпиад и конкурсов: анонсы, сроки, файлы и ссылки "
    "собраны в одном проверенном месте.",
    "Учитель публикует один раз — объявление само появляется на главной, "
    "в календаре и у подписчиков.",
    "Ученик не ищет по чатам и доскам объявлений: предмет, тип, дата — "
    "и нужный анонс найден за 30 секунд.",
    "Ничего не теряется: архив прошлых лет, регламенты и презентации "
    "доступны в любой момент.",
    "Работает круглосуточно, бесплатно и с любого устройства: телефона, "
    "планшета, компьютера.",
], size=16, gap=16)
rect(s, 8.5, 1.9, 4.25, 3.4, fill=LIGHT, line=BORDER,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
tf = textbox(s, 8.85, 2.25, 3.55, 2.7, anchor=MSO_ANCHOR.MIDDLE)
put(tf, "«Вся информация об олимпиадах и конкурсах школы — в одном окне: "
        "всегда актуальная и проверенная»", size=16.5, bold=True, color=DEEP,
    first=True, align=PP_ALIGN.CENTER, space_after=10, italic=True)
put(tf, "Портал экономит время учителей и снимает хаос в поиске информации "
        "для учеников.", size=12.5, color=MUTED, align=PP_ALIGN.CENTER,
    space_after=0)

# ---------------------------------------------------------------- слайд 3

s = new_slide("Междисциплинарные связи", 3)
boxes = [
    ("Математика и информатика",
     "Предметные олимпиады — и техническая основа портала: база данных, сайт, рассылка."),
    ("Филологический блок",
     "Русский язык, литература, иностранный языки: конкурсы, сочинения, публикации."),
    ("Естественные науки",
     "Физика, химия, биология, география: олимпиады, конференции, эксперименты."),
    ("ИЗО, музыка, технология",
     "Творческие конкурсы и разработки учителей с файлами для скачивания."),
    ("Обществознание и ОБЖ",
     "Гражданские, военно-патриотические и краеведческие акции школы."),
    ("ИКТ и информационная безопасность",
     "Пароли, модерация публикаций, защита персональных данных, работа с почтой."),
]
for i, (title, desc) in enumerate(boxes):
    col, row = i % 2, i // 2
    x = 0.6 + col * 6.33
    y = 1.35 + row * 1.6
    rect(s, x, y, 6.0, 1.4, fill=LIGHT, line=BORDER,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = textbox(s, x + 0.25, y + 0.16, 5.5, 1.1, anchor=MSO_ANCHOR.MIDDLE)
    put(tf, title, size=14.5, bold=True, color=DEEP, first=True, space_after=4)
    put(tf, desc, size=11.5, color=TEXT, space_after=0)
rect(s, 0.6, 6.2, 12.13, 0.65, fill=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
tf = textbox(s, 0.9, 6.3, 11.5, 0.45, anchor=MSO_ANCHOR.MIDDLE)
put(tf, "Портал объединяет кабинеты: кросс-предметные мероприятия — олимпиады, "
        "недели наук, конкурсы — видят все классы сразу.",
    size=13.5, bold=True, color=WHITE, first=True, align=PP_ALIGN.CENTER,
    space_after=0)

# ---------------------------------------------------------------- слайд 4

s = new_slide("Как работает публикация: от замысла до доски объявлений", 4)
bullets(s, 0.6, 1.35, 6.6, 3.6, [
    "Учитель регистрируется и создаёт анонс: название, предмет, тип, даты, "
    "описание, ссылка и файлы до 15 МБ.",
    "Администратор проверяет публикацию — на сайте только проверенная информация.",
    "«Мой кабинет»: свои анонсы, статусы, редактирование, смена пароля.",
    "Админка: очередь модерации, пользователи, подписчики, письма-уведомления.",
], size=15.5, gap=14)
flow = ["Регистрация", "Анонс", "Модерация", "Публикация"]
for i, step in enumerate(flow):
    x = 0.6 + i * 1.66
    rect(s, x, 5.35, 1.5, 0.72, fill=LIGHT, line=BORDER,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = textbox(s, x, 5.44, 1.5, 0.55, anchor=MSO_ANCHOR.MIDDLE)
    put(tf, step, size=12, bold=True, color=DEEP, first=True,
        align=PP_ALIGN.CENTER, space_after=0)
    if i < len(flow) - 1:
        put(tf, "→", size=13, bold=True, color=BLUE, align=PP_ALIGN.CENTER,
            space_after=0)
shot(s, "home.png", 7.6, 1.45, 5.2, "Главная: карточка анонса с фильтрами и поиском")

# ---------------------------------------------------------------- слайд 5

s = new_slide("Поиск и календарь: информация находится за секунды", 5)
bullets(s, 0.6, 1.35, 6.6, 4.4, [
    "Фильтры по 17 предметам и 5 типам: олимпиада, конкурс, творческая "
    "разработка, мероприятие, другое.",
    "Поиск по названию и описанию, галочка «Только действующие» — "
    "просроченные заявки не мешают.",
    "Календарь: дата мероприятия — синим, последний день подачи заявок — красным.",
    "«Ближайшие события» — список того, что будет в ближайшие недели.",
    "Карточку любого мероприятия можно распечатать одним нажатием — "
    "для доски объявлений или раздачи.",
], size=15.5, gap=14)
shot(s, "calendar.png", 7.6, 1.45, 5.2, "Календарь: сроки и даты в месячной сетке")

# ---------------------------------------------------------------- слайд 6

s = new_slide("Уведомления, документы и безопасность", 6)
bullets(s, 0.6, 1.35, 6.6, 4.4, [
    "Подписка по e-mail: ученик получает письмо о новых одобренных анонсах, "
    "отписка — по ссылке в письме.",
    "RSS-лента /rss.xml — для читалок и встраивания на школьный сайт.",
    "Печатная карточка: логотип школы, даты, описание и файлы — "
    "готова к печати без лишних элементов.",
    "Безопасность: пароли хешируются, секретный ключ хранится вне кода, "
    "файлы проверяются, всё под контролем администратора.",
], size=15.5, gap=14)
shot(s, "print.png", 7.6, 1.45, 5.2, "Печатная карточка мероприятия")

# ---------------------------------------------------------------- слайд 7

s = new_slide("Публикация в открытом доступе", 7)
# карточка «Сайт»
rect(s, 0.6, 1.45, 7.4, 1.5, fill=LIGHT, line=BORDER,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
tf = textbox(s, 0.95, 1.65, 6.8, 1.1, anchor=MSO_ANCHOR.MIDDLE)
put(tf, "🌐  Сайт в открытом доступе — работает круглосуточно:", size=13.5,
    color=MUTED, first=True, space_after=4)
p = tf.add_paragraph()
p.space_after = Pt(0)
link_run(p, SITE, SITE, size=15.5)
# карточка «GitHub»
rect(s, 0.6, 3.15, 7.4, 1.5, fill=LIGHT, line=BORDER,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
tf = textbox(s, 0.95, 3.35, 6.8, 1.1, anchor=MSO_ANCHOR.MIDDLE)
put(tf, "💻  Исходный код открыт на GitHub (публичный репозиторий):", size=13.5,
    color=MUTED, first=True, space_after=4)
p = tf.add_paragraph()
p.space_after = Pt(0)
link_run(p, GITHUB, GITHUB, size=15.5)
# RSS и хостинг
tf = textbox(s, 0.6, 4.9, 7.4, 1.9)
put(tf, "RSS-лента: " + SITE + "/rss.xml", size=13, color=TEXT, first=True,
    space_after=8)
put(tf, "Бесплатный хостинг Replit · публичный репозиторий GitHub · "
        "обновления — git pull и одна кнопка Republish.", size=13, color=MUTED,
    space_after=8)
put(tf, "Данные школы (база и файлы) в открытый репозиторий не попадают.", size=13,
    color=MUTED, space_after=0)
# QR-код
qr = qrcode.make(SITE)
buf = io.BytesIO()
qr.save(buf, format="PNG")
buf.seek(0)
rect(s, 9.05, 1.45, 3.7, 3.9, fill=WHITE, line=BORDER,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
s.shapes.add_picture(buf, Inches(9.55), Inches(1.75), width=Inches(2.7),
                     height=Inches(2.7))
tf = textbox(s, 9.3, 4.6, 3.2, 0.6)
put(tf, "Наведите камеру —", size=12, color=MUTED, first=True,
    align=PP_ALIGN.CENTER, space_after=2)
put(tf, "сразу откроется сайт портала", size=12, color=MUTED,
    align=PP_ALIGN.CENTER, space_after=0)
rect(s, 9.05, 5.55, 3.7, 1.3, fill=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
tf = textbox(s, 9.3, 5.7, 3.2, 1.0, anchor=MSO_ANCHOR.MIDDLE)
put(tf, "Спасибо за внимание!", size=16, bold=True, color=WHITE, first=True,
    align=PP_ALIGN.CENTER, space_after=4)
put(tf, "Вопросы и предложения — администрации школы", size=11, color=LIGHT,
    align=PP_ALIGN.CENTER, space_after=0)

# ---------------------------------------------------------------- сохранение

prs.core_properties.title = f"Портал олимпиад и конкурсов — {SCHOOL}"
prs.core_properties.author = SCHOOL
prs.core_properties.subject = "Презентация школьного портала (7 слайдов)"
prs.save(OUT)
print(f"OK: {OUT}")
